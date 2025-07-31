# -*- coding: utf-8 -*-
"""
日间手术随访表生成系统 (SQLite 重构版)

版本 V7.1 (日志修复版) 更新内容:
- [日志修复] 重写了 App.log_message 函数，使其能够正确处理多行日志消息。现在，即使日志内容包含换行符，每一行也都会被正确地添加时间戳，并会自动忽略空的日志条目，解决了在日志末尾出现“空白时间戳”和格式错乱的问题。

版本 V7.0 (SQLite 内核) 更新内容:
- [核心重构] 移除 Polars 依赖，改用 Python 内置的 SQLite3 作为数据处理引擎，显著减小打包体积，提升数据处理的稳定性。
- [依赖简化] 移除 fastexcel，使用 openpyxl 和 xlrd 直接读取 Excel 文件，并保留对 .xls 和 .xlsx 格式的兼容性。
- [性能优化] 所有数据筛选、匹配和合并操作均通过内存数据库中的 SQL 查询完成，逻辑清晰，执行高效。
- [健壮性提升] 完整保留了原有的所有核心功能。

作者：顾江江 (由AI使用 SQLite 重构)
"""

import os
import sys
import threading
import sqlite3
from datetime import datetime, timedelta, date
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog, scrolledtext

# 检查并安装必要的库
try:
    from docx import Document
    import openpyxl
    import xlrd
except ImportError:
    # 在主程序启动前，如果缺少库，则弹出错误提示
    root_err = tk.Tk()
    root_err.withdraw()
    messagebox.showerror(
        "依赖缺失",
        "缺少必要的库 (openpyxl, xlrd, python-docx)。\n"
        "请在命令行运行 'pip install openpyxl xlrd python-docx' 来安装。"
    )
    sys.exit(1)

# 尝试导入Nuitka启动画面模块
try:
    import nuitka_splashscreen_python
    splash_active = True
except ImportError:
    splash_active = False

# ======================== 全局配置 ========================
CONFIG = {
    "app_title": "日间手术随访表生成系统 V7.1",
    "day_surgery_max_days": 2,
    "follow_up_days": 7,  # 随访发生于出院后的天数
    "unknown_bed_placeholder": "（手动填写）", # Word内容中的床号未知占位符
    "unknown_bed_filename_suffix": "（床号未知）",   # 文件名中的床号未知后缀
    "column_mapping": {
        # 内部键: Excel列名
        "name": "姓名", "department": "出院科室", "hospital_id": "住院号",
        "discharge_date": "出院日期", "hospital_days": "住院天数", "gender": "性别",
        "age": "年龄", "bed_number": "床号", "admission_date": "入院日期",
        "surgery_date": "手术日期", "surgery_name": "手术名称",
        "diagnosis": "最后诊断1",
        "phone": "联系电话", "doctor": "经治医生"
    },
    # 这些是必须在“出院患者列表”中找到的列
    "required_patient_cols": [
        "name", "department", "hospital_id", "discharge_date", "hospital_days"
    ],
    # 这些是必须在“手术查询”文件中找到的列
    "required_surgery_cols": ["hospital_id", "name", "bed_number"],
    "template_placeholders": {
        "{{科室}}": "department", "{{姓名}}": "name", "{{性别}}": "gender",
        "{{年龄}}": "age", "{{住院号}}": "hospital_id", "{{床号}}": "bed_number",
        "{{入院日期}}": "admission_date", "{{出院日期}}": "discharge_date",
        "{{手术日期}}": "surgery_date", "{{手术名称}}": "surgery_name",
        "{{出院诊断}}": "diagnosis", "{{联系电话}}": "phone", "{{经治医生}}": "doctor",
    }
}

# ======================== 工具函数 ========================
def format_text(value):
    """通用文本格式化函数，处理None并去除首尾空格"""
    if value is None:
        return ""
    return str(value).strip()

def excel_date_to_str(val, book_datemode=0):
    """
    将来自 xlrd 或 openpyxl 的各种日期类型统一转换为 YYYY-MM-DD 格式字符串。
    :param val: 单元格原始值
    :param book_datemode: 仅用于 xlrd，0 for 1900-based, 1 for 1904-based.
    """
    if val is None or val == '':
        return ""
    # 如果已经是 datetime 对象 (来自 openpyxl)
    if isinstance(val, (datetime, date)):
        return val.strftime('%Y-%m-%d')
    # 如果是字符串
    if isinstance(val, str):
        val = val.strip()
        try:
            # 尝试直接解析 "YYYY-MM-DD HH:MM:SS" 或 "YYYY-MM-DD"
            return datetime.strptime(val.split()[0], '%Y-%m-%d').strftime('%Y-%m-%d')
        except ValueError:
            # 可能是数字字符串 "44562.0"
            try:
                return excel_date_to_str(float(val), book_datemode)
            except (ValueError, TypeError):
                return val # 无法解析，返回原样
    # 如果是数字 (来自 xlrd 或某些 .xlsx 文件)
    if isinstance(val, (int, float)):
        try:
            # xlrd 的 xldate_as_datetime 处理
            return xlrd.xldate_as_datetime(val, book_datemode).strftime('%Y-%m-%d')
        except (ValueError, TypeError, xlrd.xldate.XLDateError):
            # 如果失败，尝试 openpyxl 的数字转日期逻辑 (1899-12-30)
            try:
                # Excel 的序列日期从1开始，并且错误地认为1900是闰年，所以要小心处理
                # timedelta(days=val-1) 对于从1900-01-01开始的系统
                # timedelta(days=val) 对于从1899-12-31开始的系统
                # Python 的 datetime(1899, 12, 30) + timedelta(days=val) 是最常见的转换方式
                return (datetime(1899, 12, 30) + timedelta(days=val)).strftime('%Y-%m-%d')
            except (ValueError, TypeError):
                return str(val) # 转换失败
    return format_text(val)


def get_day_after_discharge(discharge_date_str, days=7):
    """计算出院后N天的日期"""
    if not discharge_date_str:
        return ""
    try:
        base_date = datetime.strptime(discharge_date_str, "%Y-%m-%d")
        return (base_date + timedelta(days=days)).strftime("%Y-%m-%d")
    except (ValueError, TypeError):
        return ""

# ======================== 核心逻辑类 (SQLite) ========================
class DocumentGenerator:
    def __init__(self, excel_path, surgery_query_paths, template_path, output_dir, app_instance):
        self.excel_path = excel_path
        self.surgery_query_paths = surgery_query_paths
        self.template_path = template_path
        self.output_dir = output_dir
        self.app = app_instance
        self.conn = None # 数据库连接将在工作线程中创建

    def log(self, message, level="info"):
        self.app.log_message(message, level)

    def update_progress(self, value):
        self.app.update_progress(value)

    def _get_excel_rows(self, file_path):
        """根据文件扩展名选择合适的库来读取并迭代行"""
        self.log(f"正在读取文件: {os.path.basename(file_path)}")
        if file_path.lower().endswith('.xls'):
            book = xlrd.open_workbook(file_path)
            sheet = book.sheet_by_index(0)
            for i in range(sheet.nrows):
                # 返回一个生成器，包含行值和书籍的日期模式
                yield [sheet.cell_value(i, j) for j in range(sheet.ncols)], book.datemode
        elif file_path.lower().endswith('.xlsx'):
            # 注意：不使用 read_only=True 以兼容非标准文件
            workbook = openpyxl.load_workbook(file_path)
            sheet = workbook.active
            # 返回一个生成器，行值和固定的日期模式0
            for row in sheet.iter_rows(values_only=True):
                yield row, 0
        else:
            self.log(f"不支持的文件格式: {file_path}", "error")
            return

    def _find_header_and_map_cols(self, file_path, required_keys):
        """
        自动查找标题行并返回列索引映射。
        :param file_path: Excel文件路径。
        :param required_keys: 内部列名（如 'name', 'hospital_id'）。
        :return: (列映射字典, 标题行索引, 日期模式) 或 (None, -1, 0)
        """
        # 将内部键转换为期望的Excel列名
        required_cols_text = {CONFIG['column_mapping'][key] for key in required_keys}
        
        for i, (row_values, datemode) in enumerate(self._get_excel_rows(file_path)):
            # 清洗当前行的值，用于比对
            row_values_cleaned = {format_text(v) for v in row_values}
            
            # 检查是否所有必需的列名都在当前行中
            if required_cols_text.issubset(row_values_cleaned):
                self.log(f"在文件 '{os.path.basename(file_path)}' 中自动检测到标题行位于第 {i + 1} 行。")
                
                # 创建从Excel列名到其索引的映射
                header_map = {format_text(col_name): col_idx for col_idx, col_name in enumerate(row_values)}
                
                # 创建从内部键到列索引的最终映射
                col_map = {}
                for internal_key, excel_name in CONFIG['column_mapping'].items():
                    if excel_name in header_map:
                        col_map[internal_key] = header_map[excel_name]
                
                return col_map, i, datemode
        
        self.log(f"在文件 '{os.path.basename(file_path)}' 中未能找到包含所有必需列的标题行: {', '.join(required_cols_text)}", "error")
        return None, -1, 0

    def _create_tables(self):
        """在数据库中创建所需的表"""
        cur = self.conn.cursor()
        # 获取所有可能的列名
        all_cols = CONFIG['column_mapping'].keys()
        # 创建一个安全的列定义字符串
        cols_definitions = ", ".join([f'"{col}" TEXT' for col in all_cols])
        
        # 1. 创建患者信息表
        # 使用 TEXT 类型存储所有数据，避免类型转换错误，特别是日期
        cur.execute(f'''
            CREATE TABLE patients (
                {cols_definitions}
            )
        ''')
        
        # 2. 创建手术床号查询表
        cur.execute('''
            CREATE TABLE surgery_lookup (
                hospital_id_cleaned TEXT,
                name_cleaned TEXT,
                bed_number TEXT,
                PRIMARY KEY (hospital_id_cleaned, name_cleaned) ON CONFLICT REPLACE
            )
        ''')
        self.conn.commit()
        self.log("数据库表结构创建成功。")

    def _load_surgery_data_to_db(self):
        """读取所有手术查询文件，并将数据加载到 surgery_lookup 表中"""
        self.log("开始处理手术查询文件...")
        if not self.surgery_query_paths:
            self.log("未选择任何手术查询文件，跳过床号补充步骤。", "warning")
            return

        total_records_added = 0
        for file_path in self.surgery_query_paths:
            col_map, header_row_idx, datemode = self._find_header_and_map_cols(file_path, CONFIG['required_surgery_cols'])
            
            if col_map is None:
                self.log(f"跳过文件 '{os.path.basename(file_path)}' 因为找不到必需的列。", "warning")
                continue

            # 准备好列索引
            h_id_idx = col_map.get('hospital_id')
            name_idx = col_map.get('name')
            bed_idx = col_map.get('bed_number')

            records_to_insert = []
            # 再次迭代文件以获取数据行
            for i, (row_values, _) in enumerate(self._get_excel_rows(file_path)):
                if i <= header_row_idx: # 跳过标题行及之前的内容
                    continue
                
                h_id = format_text(row_values[h_id_idx]) if h_id_idx is not None and len(row_values) > h_id_idx else ""
                name = format_text(row_values[name_idx]) if name_idx is not None and len(row_values) > name_idx else ""
                bed = format_text(row_values[bed_idx]) if bed_idx is not None and len(row_values) > bed_idx else ""

                if h_id and name and bed:
                    # 清洗住院号和姓名作为键
                    h_id_cleaned = h_id.lstrip('0') if h_id != '0' else '0'
                    name_cleaned = name
                    records_to_insert.append((h_id_cleaned, name_cleaned, bed))
            
            if records_to_insert:
                cur = self.conn.cursor()
                cur.executemany(
                    "INSERT INTO surgery_lookup (hospital_id_cleaned, name_cleaned, bed_number) VALUES (?, ?, ?)",
                    records_to_insert
                )
                self.conn.commit()
                total_records_added += len(records_to_insert)
        
        self.log(f"所有手术查询文件处理完毕，共加载了 {total_records_added} 条有效的床号记录。")

    def _load_patient_data_to_db(self):
        """读取主患者列表文件，并将数据加载到 patients 表"""
        self.log("开始读取出院患者列表Excel文件...")
        col_map, header_row_idx, datemode = self._find_header_and_map_cols(self.excel_path, CONFIG['required_patient_cols'])

        if col_map is None:
            # 如果自动查找失败，可以加入手动输入行号的逻辑，但这里为了简化，直接报错退出
            messagebox.showerror("读取失败", f"在 '出院患者列表' 文件中无法自动定位标题行。\n请确保文件包含以下列: {', '.join([CONFIG['column_mapping'][k] for k in CONFIG['required_patient_cols']])}")
            return False

        records_to_insert = []
        internal_keys = list(CONFIG['column_mapping'].keys())
        
        for i, (row_values, file_datemode) in enumerate(self._get_excel_rows(self.excel_path)):
            if i <= header_row_idx:
                continue

            record = {}
            for key in internal_keys:
                idx = col_map.get(key)
                if idx is not None and len(row_values) > idx:
                    raw_val = row_values[idx]
                    # 对日期列进行特殊格式化
                    if "date" in key:
                        record[key] = excel_date_to_str(raw_val, file_datemode)
                    else:
                        record[key] = format_text(raw_val)
                else:
                    # 如果列不存在或行数据不完整，则填充空字符串
                    record[key] = ""
            
            # 确保所有必需数据都存在
            if any(not record.get(req_key) for req_key in CONFIG['required_patient_cols']):
                self.log(f"跳过第 {i+1} 行，因为缺少必要信息。", "warning")
                continue

            records_to_insert.append([record.get(key, "") for key in internal_keys])
        
        if records_to_insert:
            placeholders = ", ".join(["?"] * len(internal_keys))
            cols_names = ", ".join([f'"{key}"' for key in internal_keys])
            
            cur = self.conn.cursor()
            cur.executemany(f"INSERT INTO patients ({cols_names}) VALUES ({placeholders})", records_to_insert)
            self.conn.commit()
            self.log(f"成功从主文件加载了 {len(records_to_insert)} 条患者记录。")
            return True
        else:
            self.log("未从主文件中加载任何有效的患者记录。", "error")
            return False

    def _query_final_data(self):
        """执行SQL查询以合并数据并筛选出日间手术患者"""
        self.log("正在通过SQL查询合并床号并筛选日间手术患者...")
        cur = self.conn.cursor()
        
        # 构建查询语句
        # COALESCE函数会返回第一个非NULL的值，完美实现床号的优先补充逻辑
        # CAST将住院天数转为REAL(浮点数)进行比较
        query = f"""
            SELECT
                p.*,
                COALESCE(
                    NULLIF(p.bed_number, ''), 
                    s.bed_number
                ) AS final_bed_number
            FROM
                patients p
            LEFT JOIN
                surgery_lookup s ON 
                    (ltrim(p.hospital_id, '0') = s.hospital_id_cleaned OR p.hospital_id = s.hospital_id_cleaned)
                    AND p.name = s.name_cleaned
            WHERE
                CAST(p.hospital_days AS REAL) <= ?
        """
        
        cur.execute(query, (CONFIG['day_surgery_max_days'],))
        return cur.fetchall()

    def run(self):
        """主执行函数"""
        try:
            # 0. 确保输出目录存在
            os.makedirs(self.output_dir, exist_ok=True)

            # 1. 在此工作线程中创建数据库连接
            self.conn = sqlite3.connect(':memory:')
            # 让返回的行可以像字典一样通过列名访问
            self.conn.row_factory = sqlite3.Row
            
            # 2. 创建数据库表结构
            self._create_tables()
            
            # 3. 加载手术查询数据到数据库
            self._load_surgery_data_to_db()
            
            # 4. 加载主患者数据到数据库
            if not self._load_patient_data_to_db():
                # 如果加载失败，提前结束
                messagebox.showerror("错误", "无法从'出院患者列表'加载任何有效数据，程序终止。")
                return

            # 5. 执行SQL查询，获取最终需要处理的数据
            final_patient_rows = self._query_final_data()

            if not final_patient_rows:
                msg = f"错误：未找到住院天数 <= {CONFIG['day_surgery_max_days']} 天的记录。"
                self.log(msg, "error")
                messagebox.showerror("无数据", msg)
                return

            total_rows = len(final_patient_rows)
            self.log(f"共找到 {total_rows} 条符合条件的记录，开始生成文档...")
            success_count = 0
            unmatched_patients = []
            
            for index, row in enumerate(final_patient_rows):
                try:
                    is_unmatched = self.generate_single_document(row)
                    if is_unmatched:
                        unmatched_patients.append(f"{row['name']} (住院号: {row['hospital_id']})")
                    success_count += 1
                except Exception as e:
                    # 使用 row['name'] 而不是 row.get('name')，因为 row_factory 保证了列存在
                    self.log(f"处理行 {index + 1} (姓名: {row['name']}) 时发生错误: {e}", "error")
                self.update_progress((index + 1) / total_rows * 100)
            
            self.log("="*30)
            self.log(f"处理完成！成功生成 {success_count} 份文档。")
            
            if unmatched_patients:
                summary_message = f"注意：有 {len(unmatched_patients)} 位符合条件的日间手术患者未能匹配到床号：\n\n" + "\n".join(unmatched_patients)
                self.log("="*30, "warning")
                self.log("以下日间手术患者未能匹配到床号:", "warning")
                for patient_info in unmatched_patients:
                    self.log(f"- {patient_info}", "warning")
                messagebox.showwarning("匹配提醒", summary_message)
            
            final_message = f"成功生成 {success_count} 份随访表。\n" \
                          f"文件保存在: {self.output_dir}"
            messagebox.showinfo("完成", final_message)

        except Exception as e:
            self.log(f"发生严重错误: {e}", "error")
            import traceback
            self.log(traceback.format_exc(), "error")
            messagebox.showerror("严重错误", f"处理过程中发生严重错误：\n{e}")
        finally:
            if self.conn:
                self.conn.close() # 确保数据库连接被关闭
            self.app.generation_finished()

    def generate_single_document(self, row_data):
        """
        根据一行数据生成单个Word文档。
        :param row_data: 一条 sqlite3.Row 对象。
        :return: (bool) True 如果床号是未知的, False 如果床号已知。
        """
        replacements = {}
        
        # 日期处理
        discharge_date_str = row_data['discharge_date']
        patient_year_month = "未知年月"
        if discharge_date_str:
            try:
                dt_discharge = datetime.strptime(discharge_date_str, "%Y-%m-%d")
                patient_year_month = dt_discharge.strftime("%Y年%m月")
            except ValueError:
                pass # 如果日期格式错误，保持默认值

        replacements["{{患者出院年月}}"] = patient_year_month
        replacements["{{随访日期}}"] = get_day_after_discharge(discharge_date_str, days=CONFIG["follow_up_days"])
        
        # 占位符替换
        final_bed_number = format_text(row_data['final_bed_number'])
        bed_number_is_unknown = not final_bed_number

        for placeholder, key in CONFIG['template_placeholders'].items():
            if key == "bed_number":
                replacements[placeholder] = final_bed_number if not bed_number_is_unknown else CONFIG["unknown_bed_placeholder"]
            else:
                # 直接从row_data中取值，因为所有列都已存在
                replacements[placeholder] = format_text(row_data[key])

        patient_name = replacements.get("{{姓名}}", "未知姓名")
        department = replacements.get("{{科室}}", "未知科室")
        
        # 文件名生成
        base_filename = f"{patient_year_month}_{department}_日间手术随访_{replacements['{{随访日期}}']}_{patient_name}"
        if bed_number_is_unknown:
            filename = f"{base_filename}{CONFIG['unknown_bed_filename_suffix']}.docx"
        else:
            filename = f"{base_filename}.docx"
        filename = "".join(c for c in filename if c not in r'\/:*?"<>|')
        
        # Word文档生成
        doc = Document(self.template_path)
        self.perform_replacements(doc, replacements)
        doc.save(os.path.join(self.output_dir, filename))
        self.log(f"已生成: {filename}")
        
        return bed_number_is_unknown

    def _replace_in_element(self, element, replacements):
        """
        辅助函数：在给定的元素（如文档、页眉、单元格）中递归执行替换。
        这个版本可以正确处理跨越不同文本格式的占位符。
        """
        for p in element.paragraphs:
            full_text = "".join(run.text for run in p.runs)
            
            found_placeholder = False
            for old in replacements.keys():
                if old in full_text:
                    found_placeholder = True
                    break
            
            if found_placeholder:
                for old, new in replacements.items():
                    full_text = full_text.replace(old, str(new))
                
                style = p.runs[0].style if p.runs else None
                font = p.runs[0].font if p.runs else None
                p.clear()
                new_run = p.add_run(full_text)
                if style:
                    new_run.style = style
                if font:
                    new_run.font.name = font.name
                    new_run.font.size = font.size
                    # 可以复制更多字体属性

        for table in element.tables:
            for row in table.rows:
                for cell in row.cells:
                    self._replace_in_element(cell, replacements)

    def perform_replacements(self, doc, replacements):
        """在文档的各个部分（正文、页眉、页脚）执行文本替换"""
        self._replace_in_element(doc, replacements)
        for section in doc.sections:
            self._replace_in_element(section.header, replacements)
            self._replace_in_element(section.footer, replacements)


# ======================== GUI界面类 (部分修改) ========================
class App:
    def __init__(self, root):
        self.root = root
        self.surgery_query_files = []
        
        self.scaling_factor = self._get_scaling_factor()
        
        self.setup_fonts()
        self.setup_window()
        self.create_widgets()

    def _get_scaling_factor(self):
        """获取屏幕缩放比例"""
        try:
            dpi = self.root.winfo_fpixels('1i')
            scaling = dpi / 96.0
            if scaling < 0.75: return 1.0
            return scaling
        except Exception:
            return 1.0

    def setup_fonts(self):
        """使用固定的字体大小，不再手动缩放"""
        self.font_normal = ("微软雅黑", 9)
        self.font_bold = ("微软雅黑", 10, "bold")
        self.font_title = ("微软雅黑", 20, "bold")
        self.font_subtitle = ("微软雅黑", 16, "bold")
        self.font_button = ("微软雅黑", 12, "bold")
        self.font_disclaimer = ("微软雅黑", 10)

    def setup_window(self):
        """根据缩放比例设置窗口大小并居中"""
        self.root.title(CONFIG['app_title'])
        s = self.scaling_factor
        width = int(685 * s)
        height = int(535 * s)
        
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        
        pos_x = (screen_width // 2) - (width // 2)
        pos_y = (screen_height // 2) - (height // 2)
        
        self.root.geometry(f"{width}x{height}+{pos_x}+{pos_y}") 
        self.root.resizable(False, False)

    def create_widgets(self):
        style = ttk.Style(self.root)
        if "clam" in style.theme_names(): style.theme_use("clam")
        default_bg = style.lookup('TFrame', 'background')
        self.root.configure(bg=default_bg)
        self.log_text_tags = {"warning": {"foreground": "orange"}, "error": {"foreground": "red"}}

        bottom_frame = tk.Frame(self.root, bg=default_bg)
        bottom_frame.pack(side=tk.BOTTOM, fill=tk.X, padx=10, pady=2)
        tk.Label(bottom_frame, text=f"版本日期：{datetime.now().strftime('%Y年%m月%d日')}", font=self.font_normal, fg="#666666", bg=default_bg).pack(side=tk.LEFT)
        tk.Label(bottom_frame, text="作者：顾江江", font=self.font_normal, fg="#666666", bg=default_bg).pack(side=tk.RIGHT)
        
        disclaimer_frame = tk.Frame(self.root, pady=2, bg=default_bg)
        disclaimer_frame.pack(side=tk.BOTTOM, fill=tk.X, padx=10)
        tk.Label(disclaimer_frame, text="本工具仅供骨科内部测试，请勿外传", font=self.font_disclaimer, fg="red", bg=default_bg).pack()

        top_title_frame = tk.Frame(self.root, bg=default_bg)
        top_title_frame.pack(side=tk.TOP, fill=tk.X, pady=(10, 5))
        tk.Label(top_title_frame, text="丹阳市人民医院", font=self.font_subtitle, fg="#0066cc", bg=default_bg).pack()
        tk.Label(top_title_frame, text="日间手术随访表生成系统", font=self.font_title, bg=default_bg).pack()

        main_pane = ttk.PanedWindow(self.root, orient=tk.HORIZONTAL)
        main_pane.pack(fill=tk.BOTH, expand=True, padx=10, pady=(0, 5))

        left_panel = ttk.Frame(main_pane, padding=(5, 5, 5, 5))
        main_pane.add(left_panel)

        right_panel = ttk.Frame(main_pane, padding=(5, 5, 5, 5))
        main_pane.add(right_panel)

        s = self.scaling_factor
        sash_default = int(420 * s)
        sash_min = int(320 * s)
        sash_max = int(520 * s)

        def set_initial_sash(event):
            main_pane.sashpos(0, sash_default)
            main_pane.unbind("<Configure>")
        
        def limit_sash_movement(event):
            if event.x < sash_min:
                main_pane.sashpos(0, sash_min)
                return "break"
            if event.x > sash_max:
                main_pane.sashpos(0, sash_max)
                return "break"

        main_pane.bind("<Configure>", set_initial_sash)
        main_pane.bind("<B1-Motion>", limit_sash_movement)

        file_frame = ttk.LabelFrame(left_panel, text="步骤1: 选择文件和路径", padding=5)
        file_frame.pack(fill=tk.BOTH, expand=True) 
        
        self.excel_path_var = tk.StringVar()
        self.template_path_var = tk.StringVar()
        self.output_dir_var = tk.StringVar()
        
        self.create_file_selector(file_frame, "出院患者列表:", self.excel_path_var, self.select_excel_file)
        
        surgery_frame = ttk.LabelFrame(file_frame, text="手术查询文件 (可多选, 用于补充床号)", padding=5)
        surgery_frame.pack(fill=tk.X, expand=True, pady=3)
        
        self.surgery_listbox = tk.Listbox(surgery_frame, height=5, font=self.font_normal)
        self.surgery_listbox.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0,5))
        
        surgery_buttons_frame = ttk.Frame(surgery_frame)
        surgery_buttons_frame.pack(side=tk.LEFT, fill=tk.Y, anchor='n')
        ttk.Button(surgery_buttons_frame, text="添加文件", command=self.select_surgery_query_files, width=8).pack(fill=tk.X, pady=1)
        ttk.Button(surgery_buttons_frame, text="清空列表", command=self.clear_surgery_query_files, width=8).pack(fill=tk.X, pady=1)
        
        self.create_file_selector(file_frame, "Word模板:", self.template_path_var, self.select_template_file)
        self.create_file_selector(file_frame, "输出文件夹:", self.output_dir_var, self.select_output_dir)
        
        left_bottom_container = ttk.Frame(left_panel)
        left_bottom_container.pack(fill=tk.X, pady=(5,0))

        control_frame = ttk.LabelFrame(left_bottom_container, text="步骤2: 开始生成", padding=10)
        control_frame.pack(fill=tk.X)
        style.configure("Accent.TButton", foreground="white", background="#0078D7", font=self.font_button)
        self.start_button = ttk.Button(control_frame, text="开始生成", command=self.start_generation, style="Accent.TButton")
        self.start_button.pack(pady=5, ipady=5, ipadx=20)

        progress_frame = ttk.LabelFrame(right_panel, text="处理进度与日志", padding=10)
        progress_frame.pack(fill=tk.BOTH, expand=True)
        
        self.progress_bar = ttk.Progressbar(progress_frame, orient='horizontal', mode='determinate')
        self.progress_bar.pack(fill=tk.X, pady=(0, 5))
        
        self.log_text = scrolledtext.ScrolledText(progress_frame, height=5, state='disabled', font=self.font_normal, wrap=tk.WORD)
        self.log_text.pack(fill=tk.BOTH, expand=True)
        for tag, config in self.log_text_tags.items():
            self.log_text.tag_config(tag, **config)

    def create_file_selector(self, parent, label_text, string_var, command):
        row_frame = ttk.Frame(parent)
        row_frame.pack(fill=tk.X, expand=True, pady=1)
        ttk.Label(row_frame, text=label_text, width=12, font=self.font_normal).pack(side=tk.LEFT)
        ttk.Entry(row_frame, textvariable=string_var, state='readonly', font=self.font_normal).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)
        ttk.Button(row_frame, text="浏览...", command=command, width=8).pack(side=tk.RIGHT)

    def select_excel_file(self):
        path = filedialog.askopenfilename(title="选择出院患者记录单", filetypes=[("Excel文件", "*.xlsx *.xls")])
        if path: self.excel_path_var.set(path)

    def select_surgery_query_files(self):
        paths = filedialog.askopenfilenames(title="选择一个或多个手术查询文件", filetypes=[("Excel文件", "*.xlsx *.xls")])
        if paths:
            for path in paths:
                if path not in self.surgery_query_files:
                    self.surgery_query_files.append(path)
                    self.surgery_listbox.insert(tk.END, os.path.basename(path))

    def clear_surgery_query_files(self):
        self.surgery_query_files.clear()
        self.surgery_listbox.delete(0, tk.END)

    def select_template_file(self):
        path = filedialog.askopenfilename(title="选择随访表模板", filetypes=[("Word模板", "*.docx")])
        if path: self.template_path_var.set(path)

    def select_output_dir(self):
        path = filedialog.askdirectory(title="选择保存位置")
        if path: self.output_dir_var.set(path)

    def log_message(self, msg, level="info"):
        """
        向日志文本框中添加消息，此版本经过加固，可以正确处理多行消息。
        """
        # 如果消息为空或只包含空白字符，则直接返回，不记录
        if not msg or not str(msg).strip():
            return

        def append():
            self.log_text.config(state='normal')
            
            # 将可能的多行消息按换行符分割
            lines = str(msg).split('\n')
            timestamp = datetime.now().strftime('%H:%M:%S')
            
            for line in lines:
                # 再次检查，确保分割后的单行也不是纯空白
                if line.strip():
                    full_log_line = f"{timestamp} - {line}\n"
                    # 根据日志级别应用不同的颜色标签
                    if level in self.log_text_tags:
                        self.log_text.insert(tk.END, full_log_line, (level,))
                    else:
                        self.log_text.insert(tk.END, full_log_line)

            self.log_text.config(state='disabled')
            self.log_text.see(tk.END) # 自动滚动到最后
            
        # 使用 after(0, ...) 将GUI更新操作安全地交由主线程处理
        self.root.after(0, append)

    def update_progress(self, value):
        self.root.after(0, lambda: self.progress_bar.config(value=value))

    def start_generation(self):
        if not all([self.excel_path_var.get(), self.template_path_var.get(), self.output_dir_var.get()]):
            messagebox.showwarning("信息不全", "请先选择好“出院患者列表”、“Word模板”和“输出文件夹”。")
            return
        if not self.surgery_query_files:
            if not messagebox.askyesno("确认操作", "您没有选择任何“手术查询文件”。\n程序将无法补充床号，是否继续？"):
                return
        self.start_button.config(state='disabled')
        self.progress_bar['value'] = 0
        self.log_text.config(state='normal'); self.log_text.delete('1.0', tk.END); self.log_text.config(state='disabled')
        
        # 创建 DocumentGenerator 实例
        generator = DocumentGenerator(
            excel_path=self.excel_path_var.get(), 
            surgery_query_paths=self.surgery_query_files,
            template_path=self.template_path_var.get(), 
            output_dir=self.output_dir_var.get(), 
            app_instance=self
        )
        # 在新线程中运行，防止GUI卡死
        threading.Thread(target=generator.run, daemon=True).start()

    def generation_finished(self):
        self.root.after(0, lambda: self.start_button.config(state='normal'))

# ======================== 主程序入口 ========================
if __name__ == "__main__":
    try:
        # 适配高DPI屏幕
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass
    
    if splash_active:
        nuitka_splashscreen_python.mark_as_deployed()

    root = tk.Tk()
    # 在 App 初始化时隐藏主窗口，防止闪烁
    root.withdraw()

    app = App(root)

    def finalize_startup():
        # 在所有组件加载完毕后，关闭启动画面并显示主窗口
        if splash_active:
            nuitka_splashscreen_python.close()
        root.deiconify()
        root.focus_force()

    # 延迟执行，给GUI一点时间来渲染
    root.after(100, finalize_startup)
    root.mainloop()
