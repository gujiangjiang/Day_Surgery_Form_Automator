# -*- coding: utf-8 -*-
"""
核心业务逻辑模块。
包含 DocumentGenerator 类，负责所有数据处理和文档生成。
"""

import os
import sqlite3
import traceback
from datetime import datetime
from tkinter import messagebox

from docx import Document
import openpyxl
import xlrd

from ..config import CONFIG
from ..utils import format_text, excel_date_to_str, get_day_after_discharge

class DocumentGenerator:
    def __init__(self, excel_path, surgery_query_paths, template_path, output_dir, app_instance):
        self.excel_path = excel_path
        self.surgery_query_paths = surgery_query_paths
        self.template_path = template_path
        self.output_dir = output_dir
        self.app = app_instance
        self.conn = None # 数据库连接将在工作线程中创建

    def log(self, message, level="info"):
        self.app.log_message(message, level)

    def update_progress(self, value):
        self.app.update_progress(value)

    def _get_excel_rows(self, file_path):
        """根据文件扩展名选择合适的库来读取并迭代行"""
        self.log(f"正在读取文件: {os.path.basename(file_path)}")
        if file_path.lower().endswith('.xls'):
            book = xlrd.open_workbook(file_path)
            sheet = book.sheet_by_index(0)
            for i in range(sheet.nrows):
                # 返回一个生成器，包含行值和书籍的日期模式
                yield [sheet.cell_value(i, j) for j in range(sheet.ncols)], book.datemode
        elif file_path.lower().endswith('.xlsx'):
            # 注意：不使用 read_only=True 以兼容非标准文件
            workbook = openpyxl.load_workbook(file_path)
            sheet = workbook.active
            # 返回一个生成器，行值和固定的日期模式0
            for row in sheet.iter_rows(values_only=True):
                yield row, 0
        else:
            self.log(f"不支持的文件格式: {file_path}", "error")
            return

    def _find_header_and_map_cols(self, file_path, required_keys):
        """
        自动查找标题行并返回列索引映射。
        :param file_path: Excel文件路径。
        :param required_keys: 内部列名（如 'name', 'hospital_id'）。
        :return: (列映射字典, 标题行索引, 日期模式) 或 (None, -1, 0)
        """
        required_cols_text = {CONFIG['column_mapping'][key] for key in required_keys}
        
        for i, (row_values, datemode) in enumerate(self._get_excel_rows(file_path)):
            row_values_cleaned = {format_text(v) for v in row_values}
            
            if required_cols_text.issubset(row_values_cleaned):
                self.log(f"在文件 '{os.path.basename(file_path)}' 中自动检测到标题行位于第 {i + 1} 行。")
                header_map = {format_text(col_name): col_idx for col_idx, col_name in enumerate(row_values)}
                col_map = {}
                for internal_key, excel_name in CONFIG['column_mapping'].items():
                    if excel_name in header_map:
                        col_map[internal_key] = header_map[excel_name]
                return col_map, i, datemode
        
        self.log(f"在文件 '{os.path.basename(file_path)}' 中未能找到包含所有必需列的标题行: {', '.join(required_cols_text)}", "error")
        return None, -1, 0

    def _create_tables(self):
        """在数据库中创建所需的表"""
        cur = self.conn.cursor()
        all_cols = CONFIG['column_mapping'].keys()
        cols_definitions = ", ".join([f'"{col}" TEXT' for col in all_cols])
        
        cur.execute(f'CREATE TABLE patients ({cols_definitions})')
        cur.execute('''
            CREATE TABLE surgery_lookup (
                hospital_id_cleaned TEXT,
                name_cleaned TEXT,
                bed_number TEXT,
                PRIMARY KEY (hospital_id_cleaned, name_cleaned) ON CONFLICT REPLACE
            )
        ''')
        self.conn.commit()
        self.log("数据库表结构创建成功。")

    def _load_surgery_data_to_db(self):
        """读取所有手术查询文件，并将数据加载到 surgery_lookup 表中"""
        self.log("开始处理手术查询文件...")
        if not self.surgery_query_paths:
            self.log("未选择任何手术查询文件，跳过床号补充步骤。", "warning")
            return

        total_records_added = 0
        for file_path in self.surgery_query_paths:
            col_map, header_row_idx, datemode = self._find_header_and_map_cols(file_path, CONFIG['required_surgery_cols'])
            
            if col_map is None:
                self.log(f"跳过文件 '{os.path.basename(file_path)}' 因为找不到必需的列。", "warning")
                continue

            h_id_idx = col_map.get('hospital_id')
            name_idx = col_map.get('name')
            bed_idx = col_map.get('bed_number')

            records_to_insert = []
            for i, (row_values, _) in enumerate(self._get_excel_rows(file_path)):
                if i <= header_row_idx:
                    continue
                
                h_id = format_text(row_values[h_id_idx]) if h_id_idx is not None and len(row_values) > h_id_idx else ""
                name = format_text(row_values[name_idx]) if name_idx is not None and len(row_values) > name_idx else ""
                bed = format_text(row_values[bed_idx]) if bed_idx is not None and len(row_values) > bed_idx else ""

                if h_id and name and bed:
                    h_id_cleaned = h_id.lstrip('0') if h_id != '0' else '0'
                    records_to_insert.append((h_id_cleaned, name, bed))
            
            if records_to_insert:
                cur = self.conn.cursor()
                cur.executemany(
                    "INSERT INTO surgery_lookup (hospital_id_cleaned, name_cleaned, bed_number) VALUES (?, ?, ?)",
                    records_to_insert
                )
                self.conn.commit()
                total_records_added += len(records_to_insert)
        
        self.log(f"所有手术查询文件处理完毕，共加载了 {total_records_added} 条有效的床号记录。")

    def _load_patient_data_to_db(self):
        """读取主患者列表文件，并将数据加载到 patients 表"""
        self.log("开始读取出院患者列表Excel文件...")
        col_map, header_row_idx, datemode = self._find_header_and_map_cols(self.excel_path, CONFIG['required_patient_cols'])

        if col_map is None:
            messagebox.showerror("读取失败", f"在 '出院患者列表' 文件中无法自动定位标题行。\n请确保文件包含以下列: {', '.join([CONFIG['column_mapping'][k] for k in CONFIG['required_patient_cols']])}")
            return False
            
        if 'bed_number' not in col_map:
            self.log("警告：主Excel文件中未找到“床号”列。将尝试从手术查询文件补充。", "warning")

        records_to_insert = []
        internal_keys = list(CONFIG['column_mapping'].keys())
        
        for i, (row_values, file_datemode) in enumerate(self._get_excel_rows(self.excel_path)):
            if i <= header_row_idx:
                continue

            record = {}
            for key in internal_keys:
                idx = col_map.get(key)
                if idx is not None and len(row_values) > idx:
                    raw_val = row_values[idx]
                    if "date" in key:
                        record[key] = excel_date_to_str(raw_val, file_datemode)
                    else:
                        record[key] = format_text(raw_val)
                else:
                    record[key] = ""
            
            if any(not record.get(req_key) for req_key in CONFIG['required_patient_cols']):
                self.log(f"跳过第 {i+1} 行，因为缺少必要信息。", "warning")
                continue

            records_to_insert.append([record.get(key, "") for key in internal_keys])
        
        if records_to_insert:
            placeholders = ", ".join(["?"] * len(internal_keys))
            cols_names = ", ".join([f'"{key}"' for key in internal_keys])
            
            cur = self.conn.cursor()
            cur.executemany(f"INSERT INTO patients ({cols_names}) VALUES ({placeholders})", records_to_insert)
            self.conn.commit()
            self.log(f"成功从主文件加载了 {len(records_to_insert)} 条患者记录。")
            return True
        else:
            self.log("未从主文件中加载任何有效的患者记录。", "error")
            return False

    def _query_final_data(self):
        """执行SQL查询以合并数据并筛选出日间手术患者"""
        self.log("正在通过SQL查询合并床号并筛选日间手术患者...")
        cur = self.conn.cursor()
        
        query = f"""
            SELECT
                p.*,
                COALESCE(
                    NULLIF(p.bed_number, ''), 
                    s.bed_number
                ) AS final_bed_number
            FROM
                patients p
            LEFT JOIN
                surgery_lookup s ON 
                    (ltrim(p.hospital_id, '0') = s.hospital_id_cleaned OR p.hospital_id = s.hospital_id_cleaned)
                    AND p.name = s.name_cleaned
            WHERE
                CAST(p.hospital_days AS REAL) <= ?
        """
        
        cur.execute(query, (CONFIG['day_surgery_max_days'],))
        return cur.fetchall()

    def run(self):
        """主执行函数"""
        try:
            os.makedirs(self.output_dir, exist_ok=True)
            self.conn = sqlite3.connect(':memory:')
            self.conn.row_factory = sqlite3.Row
            
            self._create_tables()
            self._load_surgery_data_to_db()
            
            if not self._load_patient_data_to_db():
                messagebox.showerror("错误", "无法从'出院患者列表'加载任何有效数据，程序终止。")
                return

            final_patient_rows = self._query_final_data()

            if not final_patient_rows:
                msg = f"错误：未找到住院天数 <= {CONFIG['day_surgery_max_days']} 天的记录。"
                self.log(msg, "error")
                messagebox.showerror("无数据", msg)
                return

            total_rows = len(final_patient_rows)
            self.log(f"共找到 {total_rows} 条符合条件的记录，开始生成文档...")
            success_count = 0
            unmatched_patients = []
            
            for index, row in enumerate(final_patient_rows):
                try:
                    is_unmatched = self.generate_single_document(row)
                    if is_unmatched:
                        unmatched_patients.append(f"{row['name']} (住院号: {row['hospital_id']})")
                    success_count += 1
                except Exception as e:
                    self.log(f"处理行 {index + 1} (姓名: {row['name']}) 时发生错误: {e}", "error")
                self.update_progress((index + 1) / total_rows * 100)
            
            self.app.log_raw("="*30)
            self.log(f"处理完成！成功生成 {success_count} 份文档。")
            
            if unmatched_patients:
                summary_message = f"注意：有 {len(unmatched_patients)} 位符合条件的日间手术患者未能匹配到床号：\n\n" + "\n".join(unmatched_patients)
                self.app.log_raw("="*30)
                self.log("以下日间手术患者未能匹配到床号:", "warning")
                for patient_info in unmatched_patients:
                    self.log(f"- {patient_info}", "warning")
                messagebox.showwarning("匹配提醒", summary_message)
            
            final_message = f"成功生成 {success_count} 份随访表。\n" \
                          f"文件保存在: {self.output_dir}"
            messagebox.showinfo("完成", final_message)

        except Exception as e:
            self.log(f"发生严重错误: {e}", "error")
            self.log(traceback.format_exc(), "error")
            messagebox.showerror("严重错误", f"处理过程中发生严重错误：\n{e}")
        finally:
            if self.conn:
                self.conn.close()
            self.app.generation_finished()

    def generate_single_document(self, row_data):
        """
        根据一行数据生成单个Word文档。
        :param row_data: 一条 sqlite3.Row 对象。
        :return: (bool) True 如果床号是未知的, False 如果床号已知。
        """
        replacements = {}
        
        discharge_date_str = row_data['discharge_date']
        patient_year_month = "未知年月"
        if discharge_date_str:
            try:
                dt_discharge = datetime.strptime(discharge_date_str, "%Y-%m-%d")
                patient_year_month = dt_discharge.strftime("%Y年%m月")
            except ValueError:
                pass

        replacements["{{患者出院年月}}"] = patient_year_month
        replacements["{{随访日期}}"] = get_day_after_discharge(discharge_date_str, days=CONFIG["follow_up_days"])
        
        final_bed_number = format_text(row_data['final_bed_number'])
        bed_number_is_unknown = not final_bed_number

        for placeholder, key in CONFIG['template_placeholders'].items():
            if key == "bed_number":
                replacements[placeholder] = final_bed_number if not bed_number_is_unknown else CONFIG["unknown_bed_placeholder"]
            else:
                replacements[placeholder] = format_text(row_data[key])

        patient_name = replacements.get("{{姓名}}", "未知姓名")
        department = replacements.get("{{科室}}", "未知科室")
        
        base_filename = f"{patient_year_month}_{department}_日间手术随访_{replacements['{{随访日期}}']}_{patient_name}"
        if bed_number_is_unknown:
            filename = f"{base_filename}{CONFIG['unknown_bed_filename_suffix']}.docx"
        else:
            filename = f"{base_filename}.docx"
        filename = "".join(c for c in filename if c not in r'\/:*?"<>|')
        
        doc = Document(self.template_path)
        # 更新：调用新的、保留格式的替换函数
        perform_replacements(doc, replacements)
        doc.save(os.path.join(self.output_dir, filename))
        self.log(f"已生成: {filename}")
        
        return bed_number_is_unknown


# --- 以下是新的、用于替换Word占位符并保留格式的函数 ---

def replace_in_paragraph(paragraph, replacements):
    """
    在单个段落中执行占位符替换，高级方法，可以处理跨run的占位符。
    这种方法可以保留占位符的原始样式。
    """
    for old_text, new_text in replacements.items():
        # 使用 while 循环来处理同一段落中多次出现的相同占位符
        while old_text in paragraph.text:
            runs = paragraph.runs
            full_text = "".join(run.text for run in runs)

            start_index = full_text.find(old_text)
            if start_index == -1:
                break 
            end_index = start_index + len(old_text)

            start_run_index, start_run_offset = None, None
            end_run_index, end_run_offset = None, None
            current_pos = 0

            # 定位包含占位符的起始和结束 run
            for i, run in enumerate(runs):
                run_len = len(run.text)
                if start_run_index is None and current_pos <= start_index < current_pos + run_len:
                    start_run_index = i
                    start_run_offset = start_index - current_pos
                if end_run_index is None and current_pos < end_index <= current_pos + run_len:
                    end_run_index = i
                    end_run_offset = end_index - current_pos
                current_pos += run_len
                if start_run_index is not None and end_run_index is not None:
                    break

            # 对定位到的 run 执行替换操作
            if start_run_index is not None and end_run_index is not None:
                if start_run_index == end_run_index:
                    # 情况1: 占位符在单个 run 中
                    run = runs[start_run_index]
                    run.text = run.text[:start_run_offset] + str(new_text) + run.text[end_run_offset:]
                else:
                    # 情况2: 占位符跨越多个 run
                    # 替换起始 run 的内容
                    start_run = runs[start_run_index]
                    start_run.text = start_run.text[:start_run_offset] + str(new_text)

                    # 清空中间 run 的内容
                    for i in range(start_run_index + 1, end_run_index):
                        runs[i].text = ""

                    # 处理结束 run 的内容
                    end_run = runs[end_run_index]
                    end_run.text = end_run.text[end_run_offset:]

def perform_replacements(doc, replacements):
    """
    在整个Word文档（正文、表格、页眉、页脚）中执行文本替换。
    """
    # 替换正文中的段落
    for paragraph in doc.paragraphs:
        replace_in_paragraph(paragraph, replacements)

    # 替换表格中的段落
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph, replacements)

    # 替换页眉和页脚
    for section in doc.sections:
        # 页眉
        for paragraph in section.header.paragraphs:
            replace_in_paragraph(paragraph, replacements)
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        replace_in_paragraph(paragraph, replacements)
        # 页脚
        for paragraph in section.footer.paragraphs:
            replace_in_paragraph(paragraph, replacements)
        for table in section.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        replace_in_paragraph(paragraph, replacements)