# -*- coding: utf-8 -*-
"""
模块功能：负责生成最终的Word文档。
"""
from datetime import datetime
from pathlib import Path # 导入Path类
from docx import Document
from ...config import CONFIG
from ...utils import get_day_after_discharge, format_text

def generate_single_document(row_data, template_path, output_dir):
    """
    根据一行数据生成单个Word文档。
    :param row_data: 一条 sqlite3.Row 对象。
    :param template_path: 模板文件路径 (字符串或Path对象)。
    :param output_dir: 输出目录 (字符串或Path对象)。
    :return: (is_unmatched, filename) - 床号是否未知，以及生成的文件名。
    """
    # 确保 output_dir 是 Path 对象
    output_path = Path(output_dir)

    replacements = {}
    
    discharge_date_str = row_data['discharge_date']
    patient_year_month = "未知年月"
    if discharge_date_str:
        try:
            dt_discharge = datetime.strptime(discharge_date_str, "%Y-%m-%d")
            patient_year_month = dt_discharge.strftime("%Y年%m月")
        except ValueError:
            pass

    replacements["{{患者出院年月}}"] = patient_year_month
    replacements["{{随访日期}}"] = get_day_after_discharge(discharge_date_str, days=CONFIG["follow_up_days"])
    
    final_bed_number = format_text(row_data['final_bed_number'])
    bed_number_is_unknown = not final_bed_number

    for placeholder, key in CONFIG['template_placeholders'].items():
        if key == "bed_number":
            replacements[placeholder] = final_bed_number if not bed_number_is_unknown else CONFIG["unknown_bed_placeholder"]
        else:
            replacements[placeholder] = format_text(row_data[key])

    patient_name = replacements.get("{{姓名}}", "未知姓名")
    department = replacements.get("{{科室}}", "未知科室")
    
    base_filename = f"{patient_year_month}_{department}_日间手术随访_{replacements['{{随访日期}}']}_{patient_name}"
    if bed_number_is_unknown:
        filename = f"{base_filename}{CONFIG['unknown_bed_filename_suffix']}.docx"
    else:
        filename = f"{base_filename}.docx"
    
    # 移除文件名中的非法字符
    filename = "".join(c for c in filename if c not in r'\/:*?"<>|')
    
    doc = Document(template_path)
    perform_replacements(doc, replacements)
    
    # 使用 pathlib 拼接路径并保存
    full_output_path = output_path / filename
    doc.save(full_output_path)
    
    return bed_number_is_unknown, filename

def replace_in_paragraph(paragraph, replacements):
    """在单个段落中执行占位符替换，保留样式。"""
    for old_text, new_text in replacements.items():
        while old_text in paragraph.text:
            runs = paragraph.runs
            full_text = "".join(run.text for run in runs)
            start_index = full_text.find(old_text)
            if start_index == -1: break
            end_index = start_index + len(old_text)
            start_run_index, start_run_offset, end_run_index, end_run_offset = None, None, None, None
            current_pos = 0
            for i, run in enumerate(runs):
                run_len = len(run.text)
                if start_run_index is None and current_pos <= start_index < current_pos + run_len:
                    start_run_index, start_run_offset = i, start_index - current_pos
                if end_run_index is None and current_pos < end_index <= current_pos + run_len:
                    end_run_index, end_run_offset = i, end_index - current_pos
                current_pos += run_len
                if start_run_index is not None and end_run_index is not None: break
            if start_run_index is not None and end_run_index is not None:
                if start_run_index == end_run_index:
                    run = runs[start_run_index]
                    run.text = run.text[:start_run_offset] + str(new_text) + run.text[end_run_offset:]
                else:
                    runs[start_run_index].text = runs[start_run_index].text[:start_run_offset] + str(new_text)
                    for i in range(start_run_index + 1, end_run_index): runs[i].text = ""
                    runs[end_run_index].text = runs[end_run_index].text[end_run_offset:]

def perform_replacements(doc, replacements):
    """在整个Word文档（正文、表格、页眉、页脚）中执行文本替换。"""
    for paragraph in doc.paragraphs: replace_in_paragraph(paragraph, replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs: replace_in_paragraph(paragraph, replacements)
    for section in doc.sections:
        for p in section.header.paragraphs: replace_in_paragraph(p, replacements)
        for t in section.header.tables:
            for r in t.rows:
                for c in r.cells:
                    for p in c.paragraphs: replace_in_paragraph(p, replacements)
        for p in section.footer.paragraphs: replace_in_paragraph(p, replacements)
        for t in section.footer.tables:
            for r in t.rows:
                for c in r.cells:
                    for p in c.paragraphs: replace_in_paragraph(p, replacements)
